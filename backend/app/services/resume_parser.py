import json
import os
import re
from typing import Dict, Any, Tuple
from pypdf import PdfReader
import docx
from app.ai.client import ai_client
from app.ai.prompts import RESUME_PARSING_SYSTEM_PROMPT
from app.schemas.resume import ParsedResumeData
from app.core.logging_config import logger

class ResumeParserService:
    @staticmethod
    def extract_text_from_file(file_path: str, file_type: str) -> str:
        """Extract raw text from PDF or DOCX file."""
        text = ""
        ext = file_type.upper()
        
        if ext == "PDF" or file_path.lower().endswith(".pdf"):
            try:
                reader = PdfReader(file_path)
                for page in reader.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n"
            except Exception as e:
                logger.error(f"Failed to read PDF file {file_path}: {e}")
                raise ValueError(f"Could not parse PDF: {e}")
        elif ext in ["DOCX", "DOC"] or file_path.lower().endswith(".docx"):
            try:
                doc = docx.Document(file_path)
                for para in doc.paragraphs:
                    if para.text:
                        text += para.text + "\n"
                for table in doc.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            if cell.text:
                                text += cell.text + " "
                        text += "\n"
            except Exception as e:
                logger.error(f"Failed to read DOCX file {file_path}: {e}")
                raise ValueError(f"Could not parse DOCX: {e}")
        else:
            # Fallback to plain text read
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                text = f.read()

        # Sanitize whitespace
        clean_text = re.sub(r'\r\n|\r', '\n', text)
        clean_text = re.sub(r'[ \t]+', ' ', clean_text)
        clean_text = re.sub(r'\n{3,}', '\n\n', clean_text).strip()
        
        return clean_text

    @staticmethod
    async def parse_resume_content(raw_text: str) -> ParsedResumeData:
        """Use AI to extract structured resume entities."""
        if not raw_text or len(raw_text.strip()) < 10:
            return ParsedResumeData()

        user_prompt = f"Extract structured data from the following resume text:\n\n{raw_text[:12000]}"
        raw_response = await ai_client.generate_chat_completion(
            system_prompt=RESUME_PARSING_SYSTEM_PROMPT,
            user_prompt=user_prompt,
            temperature=0.1,
            response_format_json=True
        )

        try:
            parsed_dict = json.loads(raw_response)
            return ParsedResumeData(**parsed_dict)
        except Exception as e:
            logger.warning(f"Failed to parse LLM json for resume ({e}). Attempting regex fallback.")
            # Fallback heuristic parser
            return ResumeParserService._fallback_regex_parser(raw_text)

    @staticmethod
    def _fallback_regex_parser(text: str) -> ParsedResumeData:
        email_match = re.search(r'[\w\.-]+@[\w\.-]+\.\w+', text)
        phone_match = re.search(r'(\+?\d{1,3}[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}', text)
        
        # Heuristic skills list extraction
        common_skills = [
            "Python", "JavaScript", "TypeScript", "React", "Node.js", "FastAPI", "Django",
            "SQL", "PostgreSQL", "MongoDB", "Redis", "Docker", "Kubernetes", "AWS", "GCP",
            "Machine Learning", "Deep Learning", "PyTorch", "TensorFlow", "Pandas", "NumPy",
            "LLM", "LangChain", "OpenAI", "RAG", "Git", "CI/CD", "REST API", "Microservices"
        ]
        found_skills = [skill for skill in common_skills if re.search(rf'\b{re.escape(skill)}\b', text, re.IGNORECASE)]

        return ParsedResumeData(
            email=email_match.group(0) if email_match else "",
            phone=phone_match.group(0) if phone_match else "",
            summary=text[:300].strip(),
            skills=found_skills,
        )

resume_parser = ResumeParserService()
